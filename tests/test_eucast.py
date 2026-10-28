"""Parser tests on a synthetic workbook laid out like a EUCAST table.

Every agent name and value here is an invented test fixture, not EUCAST data.
"""

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.cell.rich_text import CellRichText, TextBlock  # noqa: E402
from openpyxl.cell.text import InlineFont  # noqa: E402
from openpyxl.styles import Font  # noqa: E402

from bpdrift.eucast import EucastFormatError, read_eucast_xlsx  # noqa: E402
from bpdrift.interpret import Category, interpret  # noqa: E402
from bpdrift.pin import PinError, pin_file  # noqa: E402

SUP = InlineFont(vertAlign="superscript")


def sup(text, note):
    """Cell text followed by a superscript footnote marker."""
    return CellRichText(text, TextBlock(SUP, note))


def mic_block(ws, row, drug_class, agents):
    ws.cell(row, 1, drug_class)
    ws.cell(row, 2, "MIC breakpoints\n(mg/L)")
    ws.cell(row + 1, 2, "S ≤")
    ws.cell(row + 1, 3, "R >")
    ws.cell(row + 1, 4, "ATU")
    for i, (agent, s, r, atu) in enumerate(agents):
        ws.cell(row + 2 + i, 1, agent)
        ws.cell(row + 2 + i, 2, s)
        ws.cell(row + 2 + i, 3, r)
        if atu is not None:
            ws.cell(row + 2 + i, 4, atu)
    return row + 2 + len(agents)


def build(path, version="16.1", extra_value=None, header_r="R >"):
    wb = openpyxl.Workbook()
    wb.active.title = "Content"
    ws = wb.create_sheet("Testobacterales")
    ws.cell(1, 1, "Testobacterales")
    ws.cell(1, 9, f"EUCAST Clinical Breakpoint Tables v. {version}, valid from 2026-01-01")
    end = mic_block(ws, 8, sup("Alphacillins", "1"), [
        (sup("Alphacillin iv", "1"), "8", "8", None),
        ("Alphacillin-inhibitor", sup("8", "2"), sup("8", "2"), "16"),
        ("Alphacillin oral (other indications)", sup("(0.001)", "3"), sup("(8)", "3"), None),
        ("Betacillin", 0.5, 2, None),
        ("Gammacillin", "-", "-", None),
        ("Deltacillin", "IE", "IE", None),
        ("Epsicillin (screen only)", sup("Note", "5"), sup("Note", "5"), None),
        ("Zetacillin", "NA", "NA", None),
        ("Etacillin", "IP", "IP", None),
        (CellRichText("Thetacillin", TextBlock(SUP, "2"), ", subgroup X"), "2", "4", "4"),
    ])
    block2 = mic_block(ws, end + 2, "Omegamycins", [
        ("Omegamycin", extra_value if extra_value is not None else "1", "2", None),
    ])
    ws.cell(block2 - 2, 3, header_r)  # the R header of the second block

    an = wb.create_sheet("Anaerobes")
    an.cell(1, 1, "Anaerobes").font = Font(b=True, sz=14)  # sheet title, never a subgroup
    an.cell(4, 1, "Alphoides spp.").font = Font(b=True, sz=14)
    an.cell(6, 1, "Breakpoints for Alphoides also apply to X.").font = Font(b=True, sz=8)
    end = mic_block(an, 8, "Antimicrobial agent", [("Alphacillin", "2", "2", None)])
    an.cell(end + 2, 1, "Betoides spp.").font = Font(b=True, sz=14)
    mic_block(an, end + 4, "Antimicrobial agent", [("Alphacillin", "0.5", "0.5", None)])

    pk = wb.create_sheet("PK PD breakpoints")
    pk.cell(1, 4, "EUCAST Clinical Breakpoint Tables v. 14.0")
    mic_block(pk, 6, "Penicillins", [("Alphacillin", "2", "8", None)])
    wb.save(path)
    return path


@pytest.fixture
def table(tmp_path):
    return read_eucast_xlsx(build(tmp_path / "synthetic.xlsx"))


def row(table, agent, sheet="Testobacterales"):
    (match,) = [r for r in table.for_sheet(sheet) if r.agent == agent]
    return match


def test_version_and_sheets(table):
    assert table.version == "16.1"
    assert table.sheets == ["Testobacterales", "Anaerobes"]  # PK-PD excluded by default
    assert len(table.for_sheet("Testobacterales")) == 11


def test_subgroups_keep_repeated_agents_apart(table):
    rows = table.for_sheet("Anaerobes")
    assert [(r.subgroup, r.s.value) for r in rows] == [
        ("Alphoides spp.", 2.0),
        ("Betoides spp.", 0.5),
    ]
    assert [r.group for r in rows] == ["Alphoides spp.", "Betoides spp."]


def test_sheets_without_subgroups(table):
    for r in table.for_sheet("Testobacterales"):
        assert r.subgroup == ""
        assert r.group == "Testobacterales"


def test_superscript_notes_are_split_not_glued(table):
    r = row(table, "Alphacillin-inhibitor")
    assert r.s.kind == "value" and r.s.value == 8.0 and r.s.notes == ("2",)
    assert r.r.value == 8.0
    assert r.atu == "16"
    assert row(table, "Alphacillin iv").agent_notes == ("1",)
    assert row(table, "Alphacillin iv").drug_class == "Alphacillins"


def test_superscript_inside_agent_name(table):
    r = row(table, "Thetacillin, subgroup X")
    assert r.agent_notes == ("2",)
    assert (r.s.value, r.r.value) == (2.0, 4.0)


def test_bracketed_values(table):
    r = row(table, "Alphacillin oral (other indications)")
    assert r.s.bracketed and r.s.value == 0.001 and r.s.notes == ("3",)
    assert r.r.bracketed and r.r.value == 8.0


def test_numeric_cells(table):
    r = row(table, "Betacillin")
    assert (r.s.value, r.r.value) == (0.5, 2.0)


@pytest.mark.parametrize(
    "agent, kind",
    [("Gammacillin", "none"), ("Deltacillin", "IE"), ("Epsicillin (screen only)", "note"),
     ("Zetacillin", "NA"), ("Etacillin", "IP")],
)
def test_codes(table, agent, kind):
    r = row(table, agent)
    assert r.s.kind == kind and r.r.kind == kind
    assert interpret("1", r.breakpoint()).category is Category.NO_BREAKPOINT


def test_second_block_and_cell_reference(table):
    r = row(table, "Omegamycin")
    assert r.drug_class == "Omegamycins"
    assert r.cell == f"Testobacterales!A{r.row}"


def test_breakpoint_feeds_interpret(table):
    bp = row(table, "Betacillin").breakpoint()
    assert interpret("<=0.25", bp).category is Category.S
    assert interpret("1", bp).category is Category.I
    assert interpret(">=1", bp).category is Category.UNDECIDABLE


def test_include_pkpd(tmp_path):
    t = read_eucast_xlsx(build(tmp_path / "s.xlsx"), include_pkpd=True)
    assert "PK PD breakpoints" in t.sheets
    assert t.version == "16.1"  # the PK-PD sheet's own stale title is ignored


@pytest.mark.parametrize("bad", ["8 mg", "S", "0", "1/2"])
def test_unrecognised_value_fails_loudly(tmp_path, bad):
    path = build(tmp_path / "bad.xlsx", extra_value=bad)
    with pytest.raises(EucastFormatError, match=r"Testobacterales!B\d+"):
        read_eucast_xlsx(path)


def test_blank_spacer_row_under_header_is_skipped(tmp_path):
    # v8.1 leaves an empty row between "S ≤ / R >" and the first agent
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Testobacteriaceae"
    ws.cell(1, 9, "EUCAST Clinical Breakpoint Tables v. 8.1")
    ws.cell(5, 1, "Aminoglycosides")
    ws.cell(5, 2, "MIC breakpoint (mg/L)")
    ws.cell(6, 2, "S ≤")
    ws.cell(6, 3, "R >")
    ws.cell(8, 1, "Amikacin")
    ws.cell(8, 2, "8")
    ws.cell(8, 3, "16")
    ws.cell(9, 1, "Gentamicin")
    ws.cell(9, 2, "2")
    ws.cell(9, 3, "4")
    path = tmp_path / "spacer.xlsx"
    wb.save(path)
    t = read_eucast_xlsx(path)
    assert [(r.agent, r.s.value, r.r.value) for r in t.rows] == [("Amikacin", 8.0, 16.0), ("Gentamicin", 2.0, 4.0)]


def test_header_with_no_rows_fails_loudly(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Testobacteriaceae"
    ws.cell(1, 9, "EUCAST Clinical Breakpoint Tables v. 8.1")
    ws.cell(5, 1, "Aminoglycosides")
    ws.cell(5, 2, "MIC breakpoint (mg/L)")
    ws.cell(6, 2, "S ≤")
    ws.cell(6, 3, "R >")
    ws.cell(12, 1, "Amikacin")  # too far below the header to be this block's data
    path = tmp_path / "empty_block.xlsx"
    wb.save(path)
    with pytest.raises(EucastFormatError, match="no agent rows"):
        read_eucast_xlsx(path)


def test_header_spacing_is_ignored(tmp_path):
    # v8.1 writes the R header as "R>" with no space
    t = read_eucast_xlsx(build(tmp_path / "nospace.xlsx", header_r="R>"))
    assert any(r.agent == "Omegamycin" for r in t.rows)


def test_unexpected_header_fails_loudly(tmp_path):
    path = build(tmp_path / "hdr.xlsx", header_r="R ≥")
    with pytest.raises(EucastFormatError, match="expected 'S ≤' / 'R >'"):
        read_eucast_xlsx(path)


def test_pin_checks_checksum_and_version(tmp_path):
    path = build(tmp_path / "p.xlsx")
    good = pin_file(path, "16.1", "2026-10-26")
    assert read_eucast_xlsx(path, pin=good).pin == good

    wrong_version = pin_file(path, "15.0", "2026-10-26")
    with pytest.raises(EucastFormatError, match="pinned as v15.0"):
        read_eucast_xlsx(path, pin=wrong_version)

    other = build(tmp_path / "other.xlsx", version="13.1")
    with pytest.raises(PinError):
        read_eucast_xlsx(other, pin=good)


def test_not_a_eucast_table(tmp_path):
    path = tmp_path / "empty.xlsx"
    openpyxl.Workbook().save(path)
    with pytest.raises(EucastFormatError):
        read_eucast_xlsx(path)


def test_column_d_is_atu_only_under_an_atu_header(tmp_path):
    path = tmp_path / "v_8.1_Breakpoint_Tables.xlsx"
    build(path, version="8.1")
    wb = openpyxl.load_workbook(path, rich_text=True)   # keep the superscript footnote runs
    for ws in wb.worksheets:
        for c in ws["D"]:
            if c.value == "ATU":
                c.value = None                      # v8.1 layout: no ATU column
                ws.cell(c.row - 1, 4, "Disk content (µg)")
    wb.save(path)
    t = read_eucast_xlsx(path)
    assert t.rows and all(r.atu == "" for r in t.rows)
